"""Revisa un PDF o un Word (.docx) y resalta los tramos que parecen IA.

El texto se divide en tramos de ~200 palabras (párrafos seguidos de la misma
sección), a cada tramo se le calcula Binoculars y se compara con umbrales
calibrados en tesis humanas (resultados/umbrales.json):

  rojo      más "predecible" que el 99 % de los tramos de tesis humanas
  amarillo  más "predecible" que el 95 %
  sin color dentro de lo normal para un texto académico humano

Se omiten portada, índices, tablas, pies de página y la bibliografía.
El resultado es un indicio estadístico, no una prueba.

Con un PDF se devuelve el mismo PDF con resaltados y una portada de resumen.
Con un .docx se devuelven dos archivos: el mismo Word con resaltado (rosado =
fuerte, amarillo = moderado), comentarios y el resumen al inicio, y un PDF con
el texto analizado resaltado.

Uso:
    python -m detector.revisar tesis.pdf [salida.pdf]
    python -m detector.revisar tesis.docx [carpeta_salida]
"""

import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
UMBRALES = RAIZ / "resultados" / "umbrales.json"
PALABRAS_TRAMO = 200
MIN_PALABRAS_TRAMO = 80
MIN_PALABRAS_BLOQUE = 12
COLORES = {"alto": (0.96, 0.42, 0.38), "medio": (1.0, 0.85, 0.3)}
NOMBRES = {"alto": "Fuerte", "medio": "Moderado", "bajo": "Sin indicios"}

_FIN_BIBLIO = re.compile(r"^\s*(referencias|bibliograf[ií]a|referencias bibliogr[aá]ficas|"
                         r"anexos?|ap[eé]ndices?)\s*$", re.I)
_INDICE = re.compile(r"(\.\s?){4,}\s*\d+\s*$|^\s*(índice|tabla de contenidos?|contenido)\s*$", re.I | re.M)


@dataclass
class Bloque:
    pagina: int
    rect: tuple
    texto: str
    seccion: str = ""
    parrafo: object = None  # párrafo de python-docx, si viene de un Word


@dataclass
class Tramo:
    bloques: list = field(default_factory=list)
    senales: dict = None
    nivel: str = "bajo"

    @property
    def texto(self):
        return " ".join(b.texto for b in self.bloques)

    @property
    def palabras(self):
        return len(self.texto.split())

    @property
    def paginas(self):
        return sorted({b.pagina + 1 for b in self.bloques if b.pagina is not None})

    @property
    def ubicacion(self):
        if self.paginas:
            return "Pág. " + ", ".join(map(str, self.paginas))
        return self.bloques[0].seccion or "Sin sección"


def es_prosa(texto):
    """Descarta tablas, índices, fórmulas y listas de números."""
    pals = texto.split()
    if len(pals) < MIN_PALABRAS_BLOQUE or _INDICE.search(texto):
        return False
    letras = sum(c.isalpha() for c in texto)
    digitos = sum(c.isdigit() for c in texto)
    return letras > 0.6 * len(texto.replace(" ", "")) and digitos < 0.15 * letras


def extraer_bloques(doc):
    bloques = []
    for n, pagina in enumerate(doc):
        alto = pagina.rect.height
        for x0, y0, x1, y1, texto, *_ in pagina.get_text("blocks", sort=True):
            texto = re.sub(r"-\n(?=[a-záéíóúñ])", "", texto)
            texto = re.sub(r"\s+", " ", texto).strip()
            if _FIN_BIBLIO.match(texto) and n > len(doc) * 0.5:
                return bloques
            # Encabezados y pies de página.
            if y1 < alto * 0.07 or y0 > alto * 0.93:
                continue
            if es_prosa(texto):
                bloques.append(Bloque(n, (x0, y0, x1, y1), texto))
            elif 0 < len(texto.split()) < MIN_PALABRAS_BLOQUE and re.search(r"[A-Za-zÁÉÍÓÚÑáéíóúñ]{3}", texto):
                bloques.append(None)  # probable título: corta el tramo
    return bloques


def agrupar(bloques):
    """Tramos de ~PALABRAS_TRAMO palabras que no cruzan títulos (None).
    Un resto corto se suma al tramo anterior de la misma sección; si la
    sección entera es corta, se descarta (poca evidencia)."""
    tramos, actual, en_seccion = [], Tramo(), 0

    def cerrar():
        nonlocal actual, en_seccion
        if actual.bloques:
            if actual.palabras >= MIN_PALABRAS_TRAMO:
                tramos.append(actual)
                en_seccion += 1
            elif en_seccion:
                tramos[-1].bloques += actual.bloques
        actual = Tramo()

    for b in bloques + [None]:
        if b is None:
            cerrar()
            en_seccion = 0
            continue
        actual.bloques.append(b)
        if actual.palabras >= PALABRAS_TRAMO:
            cerrar()
    return tramos


def clasificar(tramos, umbrales):
    for t in tramos:
        b = t.senales["binoculars"]
        if b <= umbrales["alto"]:
            t.nivel = "alto"
        elif b <= umbrales["medio"]:
            t.nivel = "medio"


def resumen(tramos):
    total = sum(t.palabras for t in tramos) or 1
    return {nivel: round(100 * sum(t.palabras for t in tramos if t.nivel == nivel) / total, 1)
            for nivel in ("alto", "medio", "bajo")}


def informe(tramos, umbrales, nombre, colores=("rojo", "amarillo")):
    """Contenido del resumen como (tipo, texto, nivel): tipo en titulo,
    subtitulo, nivel, seccion, parrafo, aviso, item."""
    pct = resumen(tramos)
    fmt = lambda v: str(v).replace(".", ",")
    palabras = f"{sum(t.palabras for t in tramos):,}".replace(",", ".")
    out = [("titulo", "Revisión de posibles tramos generados por IA", None),
           ("subtitulo", nombre, None),
           ("parrafo", f"Texto analizado: {palabras} palabras en {len(tramos)} tramos "
                       "(se omiten portada, índices, tablas, títulos y bibliografía).", None),
           ("nivel", f"Indicio fuerte ({colores[0]}): {fmt(pct['alto'])} % del texto", "alto"),
           ("nivel", f"Indicio moderado ({colores[1]}): {fmt(pct['medio'])} % del texto", "medio"),
           ("nivel", f"Sin indicios: {fmt(pct['bajo'])} % del texto", "bajo"),
           ("seccion", "Cómo leer este informe", None),
           ("parrafo", "Cada tramo de unas 200 palabras se comparó con más de 300 resúmenes de tesis de "
                       f"magíster escritos por personas antes de 2022. {colores[0].capitalize()} significa que "
                       "el tramo es más predecible que el 99 % de esos textos humanos (Binoculars de "
                       f"{umbrales['alto']:.3f} o menos); {colores[1]}, que lo es más que el 95 % "
                       f"({umbrales['medio']:.3f} o menos).", None),
           ("parrafo", "En las pruebas del proyecto, con estos umbrales se marcó el "
                       f"{umbrales['sensibilidad_medio']:.0%} de los textos académicos generados por IA y el "
                       f"{umbrales['fpr_medio_prueba']:.0%} de los escritos por personas.", None),
           ("seccion", "Importante", None),
           ("aviso", "Esto es un indicio estadístico, no una prueba. Un texto humano muy formal, muy "
                     "revisado o escrito siguiendo plantillas puede salir marcado, y un texto de IA editado "
                     "a mano puede no salir. Los tramos marcados son un punto de partida para revisar, no una "
                     "conclusión. Los umbrales se calibraron con resúmenes de tesis; secciones muy "
                     "formulaicas (métodos, definiciones, marco legal) tienden a parecer más predecibles. "
                     "No debe usarse para sancionar a nadie.", None)]
    marcados = [t for t in tramos if t.nivel != "bajo"]
    if marcados:
        out.append(("seccion", "Tramos marcados", None))
        for t in marcados:
            out.append(("item", f"{t.ubicacion}: {NOMBRES[t.nivel]} (Binoculars "
                                f"{t.senales['binoculars']:.3f}): «{' '.join(t.texto.split()[:14])}...»",
                        t.nivel))
    return out


def portada(doc, tramos, umbrales, nombre):
    import pymupdf
    pagina = doc.new_page(0, width=595, height=842)
    y = 60

    def linea(texto, tam=10.5, negrita=False, color=(0.1, 0.1, 0.1), sangria=0):
        nonlocal y
        fuente = "hebo" if negrita else "helv"
        caja = pymupdf.Rect(56 + sangria, y, 540, y + 400)
        sobra = pagina.insert_textbox(caja, texto, fontsize=tam, fontname=fuente, color=color)
        y += 400 - sobra + tam * 0.5

    for tipo, texto, nivel in informe(tramos, umbrales, nombre):
        if y > 780:
            break
        if tipo == "titulo":
            linea(texto, 17, True)
        elif tipo == "subtitulo":
            linea(texto, 10, color=(0.35, 0.35, 0.35))
            y += 6
        elif tipo == "nivel":
            if nivel in COLORES:
                pagina.draw_rect(pymupdf.Rect(56, y + 1, 68, y + 11), color=None, fill=COLORES[nivel])
            linea(texto, 11, nivel != "bajo", sangria=18)
        elif tipo == "seccion":
            y += 6
            linea(texto, 12, True)
        elif tipo == "aviso":
            linea(texto, color=(0.45, 0.1, 0.1))
        elif tipo == "item":
            linea(texto, 9.5, sangria=10)
        else:
            linea(texto)


def resaltar(doc, tramos):
    for t in tramos:
        if t.nivel == "bajo":
            continue
        for b in t.bloques:
            pagina = doc[b.pagina]  # mantener la referencia mientras se edita
            anot = pagina.add_highlight_annot(_lineas(pagina, b.rect))
            anot.set_colors(stroke=COLORES[t.nivel])
            anot.set_info(title="Detector de IA",
                          content=f"{NOMBRES[t.nivel]} · Binoculars {t.senales['binoculars']:.3f}")
            anot.update()


def _lineas(pagina, rect):
    """Rectángulos de cada línea del bloque, para un resaltado limpio."""
    import pymupdf
    zona = pymupdf.Rect(rect)
    rects = []
    for bloque in pagina.get_text("dict", clip=zona)["blocks"]:
        for linea in bloque.get("lines", []):
            r = pymupdf.Rect(linea["bbox"])
            if r.intersects(zona):
                rects.append(r)
    return rects or [zona]


# --------------------------------------------------------------------- Word

_COLORES_HTML = {"alto": "#f56b61", "medio": "#ffd84d"}


def extraer_bloques_docx(doc):
    """Párrafos de prosa del cuerpo; los títulos cortan tramos. Empieza en el
    primer título de nivel 1 (tras portada e índice) y termina en
    Referencias/Anexos. Las tablas no se leen."""
    bloques, seccion, empezado = [], "", False
    for p in doc.paragraphs:
        estilo = (p.style.name if p.style is not None else "").lower()
        texto = re.sub(r"\s+", " ", p.text).strip()
        es_titulo = estilo.startswith(("heading", "título", "titulo")) or estilo == "title"
        if es_titulo:
            if _FIN_BIBLIO.match(texto) and empezado:
                break
            if estilo in ("heading 1", "título 1", "titulo 1"):
                empezado = True
            seccion = texto
            bloques.append(None)
            continue
        if not empezado or "toc" in estilo or "center" in estilo or "caption" in estilo:
            continue
        if es_prosa(texto):
            bloques.append(Bloque(None, None, texto, seccion, p))
    if not empezado:  # documento sin estilos de título: se lee todo
        bloques = [Bloque(None, None, re.sub(r"\s+", " ", p.text).strip(), "", p)
                   for p in doc.paragraphs if es_prosa(re.sub(r"\s+", " ", p.text).strip())]
    return bloques


def resaltar_docx(doc, tramos):
    from docx.enum.text import WD_COLOR_INDEX
    color = {"alto": WD_COLOR_INDEX.PINK, "medio": WD_COLOR_INDEX.YELLOW}
    for t in tramos:
        if t.nivel == "bajo":
            continue
        for i, b in enumerate(t.bloques):
            runs = [r for r in b.parrafo.runs if r.text.strip()]
            for r in runs:
                r.font.highlight_color = color[t.nivel]
            if i == 0 and runs and hasattr(doc, "add_comment"):
                doc.add_comment(runs, author="Detector de IA", initials="IA",
                                text=f"Indicio {NOMBRES[t.nivel].lower()} de IA en este tramo "
                                     f"(Binoculars {t.senales['binoculars']:.3f}). Es un indicio, no una prueba.")


def resumen_docx(doc, tramos, umbrales, nombre):
    """Inserta el resumen antes del primer párrafo, seguido de un salto de página."""
    from docx.enum.text import WD_BREAK, WD_COLOR_INDEX
    from docx.shared import Pt, RGBColor
    primero = doc.paragraphs[0]
    for tipo, texto, nivel in informe(tramos, umbrales, nombre, colores=("rosado", "amarillo")):
        p = primero.insert_paragraph_before()
        r = p.add_run(texto)
        r.font.size = Pt({"titulo": 16, "seccion": 12.5}.get(tipo, 10.5))
        r.bold = tipo in ("titulo", "seccion") or (tipo == "nivel" and nivel != "bajo")
        if tipo == "aviso":
            r.font.color.rgb = RGBColor(0x7a, 0x1a, 0x1a)
        if tipo in ("nivel", "item") and nivel in ("alto", "medio"):
            r.font.highlight_color = WD_COLOR_INDEX.PINK if nivel == "alto" else WD_COLOR_INDEX.YELLOW
    primero.insert_paragraph_before().add_run().add_break(WD_BREAK.PAGE)


def pdf_desde_docx(doc, tramos, umbrales, nombre, salida):
    """PDF con el resumen y el cuerpo analizado (títulos y párrafos de prosa),
    con los tramos marcados resaltados. No reproduce el diseño del Word."""
    import html as H
    import pymupdf
    # python-docx crea objetos nuevos en cada acceso: se compara el XML (_p).
    nivel_de = {b.parrafo._p: t for t in tramos for b in t.bloques}
    partes = []
    for tipo, texto, nivel in informe(tramos, umbrales, nombre):
        t = H.escape(texto)
        if tipo == "titulo":
            partes.append(f"<h1>{t}</h1>")
        elif tipo == "seccion":
            partes.append(f"<h3>{t}</h3>")
        elif tipo in ("nivel", "item") and nivel in _COLORES_HTML:
            partes.append(f'<p><span style="background-color:{_COLORES_HTML[nivel]}">{t}</span></p>')
        elif tipo == "aviso":
            partes.append(f'<p style="color:#7a1a1a">{t}</p>')
        elif tipo == "subtitulo":
            partes.append(f'<p style="color:#666">{t}</p>')
        else:
            partes.append(f"<p>{t}</p>")
    partes.append('<h2 style="page-break-before:always">Texto analizado</h2>'
                  '<p style="color:#666">Solo se muestran los títulos y los párrafos analizados.</p>')
    empezado = False
    for p in doc.paragraphs:
        estilo = (p.style.name if p.style is not None else "").lower()
        texto = H.escape(re.sub(r"\s+", " ", p.text).strip())
        if not texto:
            continue
        if estilo.startswith(("heading", "título", "titulo")):
            if _FIN_BIBLIO.match(p.text) and empezado:
                break
            empezado = True
            partes.append(f"<h3>{texto}</h3>")
        elif p._p in nivel_de:
            t = nivel_de[p._p]
            if t.nivel in _COLORES_HTML:
                partes.append(f'<p><span style="background-color:{_COLORES_HTML[t.nivel]}">{texto}</span></p>')
            else:
                partes.append(f"<p>{texto}</p>")
    css = ("body{font-family:sans-serif;font-size:10.5pt;line-height:1.4} h1{font-size:16pt} "
           "h2{font-size:14pt} h3{font-size:11.5pt;margin-top:10pt} p{margin:0 0 6pt 0;text-align:justify}")
    story = pymupdf.Story(html="".join(partes), user_css=css)
    writer = pymupdf.DocumentWriter(str(salida))
    pagina = pymupdf.paper_rect("a4")
    mas = 1
    while mas:
        dev = writer.begin_page(pagina)
        mas, _ = story.place(pagina + (56, 56, -56, -56))
        story.draw(dev)
        writer.end_page()
    writer.close()


# ------------------------------------------------------------------ común

def evaluar_tramos(tramos, umbrales, evaluador=None):
    if not tramos:
        raise ValueError("No se encontró texto para analizar (¿es un PDF escaneado?).")
    if evaluador is None:
        from detector.perplejidad import Evaluador
        evaluador = Evaluador(umbrales["observador"], umbrales["ejecutor"])
    for i, t in enumerate(tramos, 1):
        t.senales = evaluador.analizar(t.texto)
        print(f"\r{i}/{len(tramos)}", end="", file=sys.stderr)
    print(file=sys.stderr)
    clasificar(tramos, umbrales)


def revisar(entrada, salida=None, evaluador=None, umbrales=None):
    """Devuelve (lista de archivos de salida, tramos)."""
    entrada = Path(entrada)
    umbrales = umbrales or json.loads(UMBRALES.read_text())
    if entrada.suffix.lower() == ".docx":
        import docx
        carpeta = Path(salida) if salida else entrada.parent
        doc = docx.Document(entrada)
        tramos = agrupar(extraer_bloques_docx(doc))
        evaluar_tramos(tramos, umbrales, evaluador)
        pdf = carpeta / (entrada.stem + "_revisado.pdf")
        pdf_desde_docx(doc, tramos, umbrales, entrada.name, pdf)
        resaltar_docx(doc, tramos)
        resumen_docx(doc, tramos, umbrales, entrada.name)
        word = carpeta / (entrada.stem + "_revisado.docx")
        doc.save(word)
        return [word, pdf], tramos
    import pymupdf
    salida = Path(salida) if salida else entrada.with_name(entrada.stem + "_revisado.pdf")
    doc = pymupdf.open(entrada)
    tramos = agrupar(extraer_bloques(doc))
    evaluar_tramos(tramos, umbrales, evaluador)
    resaltar(doc, tramos)
    portada(doc, tramos, umbrales, entrada.name)
    doc.save(salida, garbage=3, deflate=True)
    return [salida], tramos


def main(argv):
    salidas, tramos = revisar(argv[1], argv[2] if len(argv) > 2 else None)
    print(json.dumps({"salida": [str(s) for s in salidas], "porcentajes": resumen(tramos),
                      "tramos": [{"ubicacion": t.ubicacion, "palabras": t.palabras, "nivel": t.nivel,
                                  "binoculars": t.senales["binoculars"]} for t in tramos]},
                     ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main(sys.argv)
