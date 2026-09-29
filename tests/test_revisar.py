import tempfile
import unittest
from pathlib import Path

from detector.revisar import Bloque, agrupar, clasificar, es_prosa, resumen

PARRAFO = ("La investigación analiza el efecto de la política pública sobre la calidad de la "
           "educación en las escuelas municipales de la región durante la última década. ")


def bloque(palabras, pagina=0):
    texto = " ".join((PARRAFO * 20).split()[:palabras])
    return Bloque(pagina, (72, 100, 500, 200), texto)


class FakeEvaluador:
    """Binoculars bajo (parece IA) si el texto menciona 'predecible'."""

    def analizar(self, texto):
        return {"binoculars": 0.7 if "predecible" in texto else 1.0}


class TestRevisar(unittest.TestCase):
    def test_es_prosa(self):
        self.assertTrue(es_prosa(PARRAFO))
        self.assertFalse(es_prosa("Capítulo 1"))
        self.assertFalse(es_prosa("Introducción ........................ 12"))
        self.assertFalse(es_prosa("12,5 13,1 14,8 15,0 16,2 17,9 18,3 19,4 20,1 21,7 22,2 23,8"))

    def test_agrupar_no_cruza_titulos(self):
        tramos = agrupar([bloque(150), bloque(120), None, bloque(90), None, bloque(30)])
        self.assertEqual([t.palabras for t in tramos], [270, 90])

    def test_agrupar_resto_corto_se_une(self):
        tramos = agrupar([bloque(200), bloque(40)])
        self.assertEqual([t.palabras for t in tramos], [240])

    def test_clasificar_y_resumen(self):
        tramos = agrupar([bloque(200), None, bloque(100)])
        tramos[0].senales = {"binoculars": 0.80}
        tramos[1].senales = {"binoculars": 0.95}
        clasificar(tramos, {"alto": 0.82, "medio": 0.88})
        self.assertEqual([t.nivel for t in tramos], ["alto", "bajo"])
        self.assertEqual(resumen(tramos), {"alto": 66.7, "medio": 0.0, "bajo": 33.3})

    def test_pdf_completo(self):
        import pymupdf
        from detector.revisar import revisar
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "t.pdf"
            doc = pymupdf.open()
            for i, extra in enumerate(["", " Texto muy predecible."]):
                p = doc.new_page()
                p.insert_textbox(pymupdf.Rect(72, 72, 523, 100), f"Capítulo {i + 1}", fontsize=14)
                p.insert_textbox(pymupdf.Rect(72, 120, 523, 780), PARRAFO * 8 + extra, fontsize=11)
            doc.save(entrada)
            umbrales = {"alto": 0.8, "medio": 0.9, "sensibilidad_medio": 0.5, "fpr_medio_prueba": 0.05}
            (salida,), tramos = revisar(entrada, evaluador=FakeEvaluador(), umbrales=umbrales)
            self.assertEqual([t.nivel for t in tramos], ["bajo", "alto"])
            out = pymupdf.open(salida)
            self.assertEqual(len(out), 3)  # portada + 2 páginas
            self.assertEqual(len(list(out[2].annots())), 1)
            self.assertEqual(len(list(out[1].annots())), 0)

    def test_docx_completo(self):
        import docx
        from detector.revisar import revisar
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "t.docx"
            doc = docx.Document()
            doc.add_paragraph("Portada de la tesis", style="Title")
            doc.add_heading("CAPÍTULO I", level=1)
            doc.add_paragraph(PARRAFO * 8)
            doc.add_heading("1.1. Sección", level=2)
            doc.add_paragraph(PARRAFO * 8 + " Texto muy predecible.")
            doc.add_heading("REFERENCIAS", level=1)
            doc.add_paragraph("Pérez, J. (2010). " + PARRAFO * 3)
            doc.save(entrada)
            umbrales = {"alto": 0.8, "medio": 0.9, "sensibilidad_medio": 0.5, "fpr_medio_prueba": 0.05}
            (word, pdf), tramos = revisar(entrada, evaluador=FakeEvaluador(), umbrales=umbrales)
            self.assertEqual([t.nivel for t in tramos], ["bajo", "alto"])
            self.assertEqual(tramos[1].ubicacion, "1.1. Sección")
            out = docx.Document(word)
            marcados = [p for p in out.paragraphs if any(r.font.highlight_color for r in p.runs)]
            self.assertTrue(out.paragraphs[0].text.startswith("Revisión de posibles tramos"))
            self.assertTrue(any("predecible" in p.text for p in marcados))
            import pymupdf
            texto_pdf = " ".join("".join(pg.get_text() for pg in pymupdf.open(pdf)).split())
            self.assertIn("Texto muy predecible", texto_pdf)
            self.assertNotIn("Pérez, J.", texto_pdf)


if __name__ == "__main__":
    unittest.main()
